from io import BytesIO

import pytest
from django.contrib.auth.models import Group, Permission, User
from django.urls import reverse
from openpyxl import load_workbook

from fretes.models import Cliente, Garantia, Produto


@pytest.fixture(autouse=True)
def staticfiles_sem_manifest(settings):
    settings.STORAGES = {
        **settings.STORAGES,
        "staticfiles": {"BACKEND": "django.contrib.staticfiles.storage.StaticFilesStorage"},
    }


@pytest.fixture
def garantia_registrada(db):
    cliente = Cliente.objects.create(nome="Cliente Teste", cnpj="00.000.000/0001-00")
    produto = Produto.objects.create(codigo="PEC-001", descricao="Descricao da peca")
    Garantia.objects.create(
        cliente=cliente,
        codigo_peca=produto.codigo,
        marca="MARCA TESTE",
        quantidade=2,
        defeito="Defeito Teste",
        nota_recebida="NF-001",
        valor="150.00",
        data_recebimento="2026-08-14",
    )
    return produto


@pytest.mark.django_db
def test_apenas_grupo_garantia_recebe_permissao_gerencial():
    permission = Permission.objects.get(
        content_type__app_label="fretes",
        codename="can_view_warranty_management",
    )

    assert Group.objects.get(name="garantia").permissions.filter(pk=permission.pk).exists()
    assert not Group.objects.get(name="expedicao").permissions.filter(pk=permission.pk).exists()
    assert not Group.objects.get(name="conferencia").permissions.filter(pk=permission.pk).exists()


@pytest.mark.django_db
def test_usuario_de_garantia_acessa_gerencial_e_ve_descricao(client, garantia_registrada):
    usuario = User.objects.create_user(username="garantia", password="senha")
    usuario.groups.add(Group.objects.get(name="garantia"))
    client.force_login(usuario)

    response = client.get(reverse("fretes:garantias_gerencial"))

    assert response.status_code == 200
    assert response.context["produtos"][0]["descricao"] == garantia_registrada.descricao
    assert b"GERENCIAL" in response.content


@pytest.mark.django_db
def test_permissao_antiga_de_relatorios_nao_libera_gerencial(client):
    usuario = User.objects.create_user(username="relatorios", password="senha")
    usuario.user_permissions.add(
        Permission.objects.get(
            content_type__app_label="fretes",
            codename="can_view_reports",
        )
    )
    client.force_login(usuario)

    response = client.get(reverse("fretes:garantias_gerencial"))

    assert response.status_code == 403


@pytest.mark.django_db
def test_superusuario_acessa_gerencial(client):
    usuario = User.objects.create_superuser(
        username="administrador",
        email="admin@example.com",
        password="senha",
    )
    client.force_login(usuario)

    response = client.get(reverse("fretes:garantias_gerencial"))

    assert response.status_code == 200


@pytest.mark.django_db
def test_excel_gerencial_exibe_descricao_na_segunda_coluna(client, garantia_registrada):
    usuario = User.objects.create_user(username="garantia-excel", password="senha")
    usuario.groups.add(Group.objects.get(name="garantia"))
    client.force_login(usuario)

    response = client.get(reverse("fretes:garantias_gerencial"), {"export": "xlsx"})
    workbook = load_workbook(BytesIO(response.content))
    worksheet = workbook.active

    assert response.status_code == 200
    assert [cell.value for cell in worksheet[1]] == [
        "Cod. Peca",
        "Descricao",
        "Quantidade",
        "Valor",
    ]
    assert worksheet.cell(row=2, column=2).value == garantia_registrada.descricao


@pytest.mark.django_db
def test_pdf_gerencial_continua_sendo_gerado(client, garantia_registrada):
    usuario = User.objects.create_user(username="garantia-pdf", password="senha")
    usuario.groups.add(Group.objects.get(name="garantia"))
    client.force_login(usuario)

    response = client.get(reverse("fretes:garantias_gerencial"), {"export": "pdf"})

    assert response.status_code == 200
    assert response["Content-Type"] == "application/pdf"
    assert response.content.startswith(b"%PDF")
