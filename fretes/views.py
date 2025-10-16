from decimal import Decimal, InvalidOperation
import json
from django.shortcuts import render, redirect, get_object_or_404
from django.http import Http404
from django.forms import inlineformset_factory
from django import forms as djforms
from django.contrib.admin.views.decorators import staff_member_required
from django.contrib.auth.decorators import permission_required, login_required
from django.db import transaction
from openpyxl import load_workbook
from django.utils import timezone
from django.contrib import messages
from django.db.models import Q, Sum
from .models import Pedido, PedidoVolume, Carrier, FreteCalculado, Produto, Cliente, Garantia
from .forms import CalcularFreteForm, PedidoForm, PedidoVolumeForm, ProdutoForm, ClienteForm, GarantiaForm, GarantiaHeaderForm, GarantiaItemForm
from .filters import FreteCalculadoFilter
from .services import calcular_frete
from .exports import exportar_fretes_excel, exportar_fretes_pdf, exportar_garantias_excel, exportar_garantias_pdf, exportar_pedido_excel, exportar_pedido_pdf, exportar_produtos_excel
from django.contrib.auth import get_user_model
from .models import AuditLog
from django.utils.dateparse import parse_date

def _last_non_empty_param(request, *keys, suffixes=()):
    querydict = request.GET

    def _clean(value):
        if value is None:
            return ""
        value = str(value).strip()
        return value

    for key in keys:
        values = querydict.getlist(key)
        if not values:
            continue
        for value in reversed(values):
            cleaned = _clean(value)
            if cleaned:
                return cleaned

    if suffixes:
        ordered_keys = list(querydict.keys())
        for key in reversed(ordered_keys):
            if not any(key.endswith(suffix) for suffix in suffixes):
                continue
            values = querydict.getlist(key)
            if not values:
                continue
            for value in reversed(values):
                cleaned = _clean(value)
                if cleaned:
                    return cleaned

    return ""

VolumeFormSet = inlineformset_factory(
    Pedido, PedidoVolume,
    form=PedidoVolumeForm,
    extra=1,                 # sempre mostra 1 linha vazia
    can_delete=True)

# fretes/views.py



VolumeFormSet = inlineformset_factory(
    Pedido,
    PedidoVolume,
    form=PedidoVolumeForm,
    extra=1,
    can_delete=True
)

def _calcular_m3_total_pedido(pedido):
    """
    Funcao auxiliar para somar a cubagem de todos os volumes de um pedido,
    considerando a quantidade.
    """
    total_m3 = Decimal("0")
    for volume in pedido.volumes.all():
        largura = volume.largura_cm / Decimal("100")
        altura = volume.altura_cm / Decimal("100")
        comprimento = volume.comprimento_cm / Decimal("100")

        # Nota: multiplica o volume pela quantidade antes de somar
        m3_volume = (largura * altura * comprimento) * Decimal(volume.quantidade)
        total_m3 += m3_volume
    return total_m3.quantize(Decimal("0.001"))


def _total_volumes_from_formset(formset):
    total = 0
    for form in getattr(formset, "forms", []):
        prefix = form.prefix
        if form.is_bound:
            delete_value = form.data.get(f"{prefix}-DELETE")
        else:
            initial = getattr(form, "initial", None) or {}
            delete_value = initial.get("DELETE")
        if delete_value is True or (isinstance(delete_value, str) and delete_value.lower() in {"1", "true", "on", "yes"}):
            continue
        qty_raw = None
        if form.is_bound:
            qty_raw = form.data.get(f"{prefix}-quantidade")
        if qty_raw in (None, ""):
            initial = getattr(form, "initial", None) or {}
            qty_raw = initial.get("quantidade")
        if qty_raw in (None, "") and getattr(form, "instance", None) is not None:
            qty_raw = getattr(form.instance, "quantidade", None)
        try:
            total += int(qty_raw)
        except (TypeError, ValueError):
            continue
    return total


def _parse_decimal_value(value):
    if value is None or value == "":
        return Decimal("0")
    if isinstance(value, (int, float)):
        return Decimal(str(value))
    s = str(value).strip()
    s = s.replace(".", "").replace(",", ".")
    try:
        return Decimal(s)
    except InvalidOperation:
        return Decimal("0")


def pedido_create(request):
    if request.method == "POST":
        form = PedidoForm(request.POST)
        formset = VolumeFormSet(request.POST, prefix="vol")
        if form.is_valid() and formset.is_valid():
            numero_pedido = form.cleaned_data.get("numero_pedido")
            pedido, created = Pedido.objects.get_or_create(
                numero_pedido=numero_pedido,
                defaults={
                    "picking": form.cleaned_data.get("picking"),
                    "carrier": form.cleaned_data.get("carrier"),
                }
            )
            if not created:
                pedido.picking = form.cleaned_data.get("picking")
                pedido.carrier = form.cleaned_data.get("carrier")
                pedido.save(update_fields=["picking", "carrier"])

            # Nota: removemos a linha 'pedido.volumes.all().delete()'
            #    que estava causando o problema.

            # Associa a instancia do pedido ao formset e salva
            formset.instance = pedido
            formset.save()

            # Atualiza m3 total do pedido apos salvar volumes
            try:
                pedido.refresh_from_db()
                pedido.m3 = _calcular_m3_total_pedido(pedido)
                pedido.save(update_fields=["m3"])
            except Exception:
                pass

            messages.success(request, "Pedido salvo com sucesso.")
            return redirect("fretes:pedido_list")
        else:
            # Logica de erro para formulario
            for err in formset.non_form_errors():
                messages.error(request, err)
            for f in formset.forms:
                for field, errs in f.errors.items():
                    messages.error(request, f"Volume: {field} -> {', '.join(errs)}")
    else:
        form = PedidoForm()
        formset = VolumeFormSet(prefix="vol")

    total_volumes = _total_volumes_from_formset(formset)
    return render(
        request,
        "fretes/pedidos_form.html",
        {"form": form, "formset": formset, "is_new": True, "total_volumes": total_volumes},
    )
# fretes/views.py



def index_view(request):
    return render(request, "fretes/index.html")

def pedido_list(request):
    q = request.GET.get("q", "").strip()
    pedidos = (
        Pedido.objects.select_related("carrier")
        .prefetch_related("volumes")
        .annotate(total_volumes=Sum("volumes__quantidade"))
        .order_by("-id")
    )
    if q:
        pedidos = pedidos.filter(numero_pedido__icontains=q)
    count = pedidos.count()
    return render(
        request,
        "fretes/pedidos_list.html",
        {"pedidos": pedidos, "q": q, "count": count, "has_filter": bool(q)},
    )


    
def pedido_update(request, pk: int):
    pedido = get_object_or_404(Pedido, pk=pk)
    if request.method == "POST":
        form = PedidoForm(request.POST, instance=pedido)
        formset = VolumeFormSet(request.POST, instance=pedido, prefix="vol")  # << prefix
        if form.is_valid() and formset.is_valid():
            form.save()
            formset.save()
            # Recalcula m3 apos atualizacao
            pedido.refresh_from_db()
            pedido.m3 = _calcular_m3_total_pedido(pedido)
            pedido.save(update_fields=["m3"])
            messages.success(request, "Pedido atualizado.")
            return redirect("fretes:pedido_list")
        else:
            for err in formset.non_form_errors():
                messages.error(request, err)
            for f in formset.forms:
                for field, errs in f.errors.items():
                    messages.error(request, f"Volume: {field} -> {', '.join(errs)}")
    else:
        form = PedidoForm(instance=pedido)
        formset = VolumeFormSet(instance=pedido, prefix="vol")  # << prefix

    total_volumes = _total_volumes_from_formset(formset)
    return render(
        request,
        "fretes/pedidos_form.html",
        {
            "form": form,
            "formset": formset,
            "pedido": pedido,
            "is_new": False,
            "total_volumes": total_volumes,
        },
    )

@permission_required('fretes.view_pedido', raise_exception=True)
def pedido_relatorio(request, pk: int):
    formato = (request.GET.get("format") or "pdf").lower()
    queryset = Pedido.objects.select_related("carrier").prefetch_related("volumes")
    pedido = get_object_or_404(queryset, pk=pk)
    if formato == "xlsx":
        return exportar_pedido_excel(pedido)
    if formato == "pdf":
        return exportar_pedido_pdf(pedido)
    raise Http404("Formato nao suportado")

@permission_required('fretes.can_view_reports', raise_exception=True)
def relatorios_view(request):
    base_qs = FreteCalculado.objects.select_related("carrier").prefetch_related("pedidos").all()
    f = FreteCalculadoFilter(request.GET, queryset=base_qs)
    export = request.GET.get("export")
    if export == "xlsx":
        return exportar_fretes_excel(f.qs)
    if export == "pdf":
        return exportar_fretes_pdf(f.qs)
    return render(request, "fretes/relatorios.html", {"filter": f})

@permission_required('fretes.can_view_reports', raise_exception=True)
def relatorios_pdf_view(request):
    f = FreteCalculadoFilter(request.GET, queryset=FreteCalculado.objects.select_related("carrier").prefetch_related("pedidos").all())
    return exportar_fretes_pdf(f.qs)

def pedidos_autocomplete(request):
    q = request.GET.get("numero_pedido", "").strip()
    qs = Pedido.objects.all().order_by("-id")
    if q:
        qs = qs.filter(numero_pedido__icontains=q)
    qs = qs[:20]
    return render(request, "fretes/_pedidos_datalist.html", {"qs": qs})


# Produtos
def produto_list(request):
    codigo = request.GET.get("codigo", "").strip()
    descricao = request.GET.get("descricao", "").strip()
    qs = Produto.objects.all().order_by("codigo")
    if codigo:
        qs = qs.filter(codigo__icontains=codigo)
    if descricao:
        qs = qs.filter(descricao__icontains=descricao)
    if request.GET.get("export") == "xlsx":
        return exportar_produtos_excel(qs)
    return render(request, "fretes/produtos_list.html", {"produtos": qs, "codigo": codigo, "descricao": descricao})


@permission_required('fretes.add_produto', raise_exception=True)
def produto_create(request):
    if request.method == "POST":
        form = ProdutoForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, "Produto cadastrado com sucesso.")
            return redirect("fretes:produto_list")
    else:
        form = ProdutoForm()
    return render(request, "fretes/produtos_form.html", {"form": form, "is_new": True})


@permission_required('fretes.change_produto', raise_exception=True)
def produto_update(request, pk: int):
    produto = get_object_or_404(Produto, pk=pk)
    if request.method == "POST":
        form = ProdutoForm(request.POST, instance=produto)
        if form.is_valid():
            form.save()
            messages.success(request, "Produto atualizado.")
            return redirect("fretes:produto_list")
    else:
        form = ProdutoForm(instance=produto)
    return render(request, "fretes/produtos_form.html", {"form": form, "is_new": False, "produto": produto})


def produtos_autocomplete(request):
    q = (request.GET.get("codigo") or request.GET.get("codigo_peca") or "").strip()
    qs = Produto.objects.all().order_by("codigo")
    if q:
        qs = qs.filter(codigo__icontains=q)
    qs = qs[:20]
    return render(request, "fretes/_produtos_datalist.html", {"qs": qs})


def produto_descricao_fragment(request):
    codigo = (request.GET.get("codigo") or request.GET.get("codigo_peca") or "").strip()
    desc = ""
    if codigo:
        p = Produto.objects.filter(codigo=codigo).first()
        if p:
            desc = p.descricao
    return render(request, "fretes/_produto_descricao.html", {"descricao": desc})


# Clientes
def cliente_list(request):
    cnpj = request.GET.get("cnpj", "").strip()
    nome = request.GET.get("nome", "").strip()
    qs = Cliente.objects.all().order_by("nome")
    if cnpj:
        qs = qs.filter(cnpj__icontains=cnpj)
    if nome:
        qs = qs.filter(nome__icontains=nome)
    return render(request, "fretes/clientes_list.html", {"clientes": qs, "cnpj": cnpj, "nome": nome})


@permission_required('fretes.add_cliente', raise_exception=True)
def cliente_create(request):
    if request.method == "POST":
        form = ClienteForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, "Cliente cadastrado com sucesso.")
            return redirect("fretes:cliente_list")
    else:
        form = ClienteForm()
    return render(request, "fretes/clientes_form.html", {"form": form, "is_new": True})


@permission_required('fretes.change_cliente', raise_exception=True)
def cliente_update(request, pk: int):
    cliente = get_object_or_404(Cliente, pk=pk)
    if request.method == "POST":
        form = ClienteForm(request.POST, instance=cliente)
        if form.is_valid():
            form.save()
            messages.success(request, "Cliente atualizado.")
            return redirect("fretes:cliente_list")
    else:
        form = ClienteForm(instance=cliente)
    return render(request, "fretes/clientes_form.html", {"form": form, "is_new": False, "cliente": cliente})


# Garantias
@permission_required('fretes.view_garantia', raise_exception=True)
def garantia_list(request):
    nota = request.GET.get("nota", "").strip()
    cnpj = request.GET.get("cnpj", "").strip()
    nome = request.GET.get("nome", "").strip()
    start_date = request.GET.get("start_date", "").strip()
    end_date = request.GET.get("end_date", "").strip()
    status = request.GET.get("status", "").strip()  # em_aberto | atendido

    qs = Garantia.objects.select_related("cliente").all().order_by("-data_recebimento", "-id")
    if nota:
        qs = qs.filter(Q(nota_recebida__icontains=nota) | Q(nota_retorno__icontains=nota))
    if cnpj:
        qs = qs.filter(cliente__cnpj__icontains=cnpj)
    if nome:
        qs = qs.filter(cliente__nome__icontains=nome)
    if start_date:
        qs = qs.filter(data_recebimento__gte=start_date)
    if end_date:
        qs = qs.filter(data_recebimento__lte=end_date)
    if status == Garantia.STATUS_EM_ABERTO:
        qs = qs.filter(Q(nota_retorno__isnull=True) | Q(nota_retorno__exact=""))
    elif status == Garantia.STATUS_ATENDIDO:
        qs = qs.filter(nota_retorno__isnull=False).exclude(nota_retorno__exact="")
    export = request.GET.get("export")
    if export == "xlsx":
        return exportar_garantias_excel(qs)
    if export == "pdf":
        return exportar_garantias_pdf(qs)

    context = {
        "garantias": qs,
        "nota": nota,
        "cnpj": cnpj,
        "nome": nome,
        "start_date": start_date,
        "end_date": end_date,
        "status": status,
    }
    return render(request, "fretes/garantias_list.html", context)


def garantia_create(request):
    if request.method == "POST":
        form = GarantiaForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, "Garantia registrada com sucesso.")
            return redirect("fretes:garantia_list")
        else:
            messages.error(request, "Corrija os erros do formulário.")
    else:
        form = GarantiaForm()
    produtos = Produto.objects.all().order_by("codigo")
    return render(request, "fretes/garantias_form.html", {"form": form, "is_new": True, "produtos": produtos})


def garantia_update(request, pk: int):
    garantia = get_object_or_404(Garantia, pk=pk)
    if request.method == "POST":
        form = GarantiaForm(request.POST, instance=garantia)
        if form.is_valid():
            form.save()
            messages.success(request, "Garantia atualizada.")
            return redirect("fretes:garantia_list")
        else:
            messages.error(request, "Corrija os erros do formulário.")
    else:
        form = GarantiaForm(instance=garantia)
    produtos = Produto.objects.all().order_by("codigo")
    return render(request, "fretes/garantias_form2.html", {"form": form, "is_new": False, "garantia": garantia, "produtos": produtos})


def produtos_options(request):
    q = _last_non_empty_param(request, "q", "codigo", "codigo_peca", suffixes=("-q",))
    qs = Produto.objects.all().order_by("codigo")
    if q:
        qs = qs.filter(Q(codigo__icontains=q) | Q(descricao__icontains=q))
    value = _last_non_empty_param(request, "codigo_peca", suffixes=("-codigo_peca",))
    return render(request, "fretes/_produto_options.html", {"qs": qs, "value": value})

def clientes_options(request):
    q = _last_non_empty_param(request, "q", "cnpj", "nome")
    qs = Cliente.objects.all().order_by("nome")
    if q:
        qs = qs.filter(Q(cnpj__icontains=q) | Q(nome__icontains=q))
    value = request.GET.get("cliente", "")
    return render(request, "fretes/_cliente_options.html", {"qs": qs, "value": value})


@permission_required('fretes.add_garantia', raise_exception=True)
def garantia_create_multi(request):
    produtos_qs = Produto.objects.all().order_by("codigo")
    produto_choices = [(p.codigo, f"{p.codigo} - {p.descricao}") for p in produtos_qs]

    if request.method == "POST":
        header_form = GarantiaHeaderForm(request.POST)
        raw_items = request.POST.get("items_payload") or "[]"
        try:
            items_data = json.loads(raw_items)
            if not isinstance(items_data, list):
                raise ValueError
        except (json.JSONDecodeError, ValueError):
            items_data = []
            header_form.add_error(None, "Nao foi possivel interpretar a lista de produtos enviada.")

        cleaned_items = []
        for idx, data in enumerate(items_data, start=1):
            item_form = GarantiaItemForm(data)
            item_form.fields["codigo_peca"].choices = produto_choices
            if item_form.is_valid():
                cleaned_items.append(item_form.cleaned_data)
            else:
                for field, errors in item_form.errors.items():
                    messages.error(
                        request,
                        f"Item {idx}: {field} - {', '.join(errors)}",
                    )

        if not cleaned_items:
            messages.error(request, "Adicione ao menos um produto antes de salvar.")

        if header_form.is_valid() and cleaned_items:
            dados = header_form.cleaned_data
            created = 0
            with transaction.atomic():
                for item in cleaned_items:
                    Garantia.objects.create(
                        cliente=dados["cliente"],
                        codigo_peca=item["codigo_peca"],
                        marca=item.get("marca") or "Nao Informado",
                        quantidade=1,
                        defeito=item["defeito"],
                        numero_lote=item.get("numero_lote", ""),
                        nota_recebida=dados["nota_recebida"],
                        valor=item.get("valor") or 0,
                        data_recebimento=dados["data_recebimento"],
                        nota_retorno=item.get("nota_retorno", ""),
                        data_retorno=item.get("data_retorno"),
                        mao_de_obra=item.get("mao_de_obra") or False,
                        valor_mao_de_obra=item.get("valor_mao_de_obra") or 0,
                    )
                    created += 1
            messages.success(request, f"{created} produto(s) adicionados a garantia.")
            return redirect("fretes:garantia_list")

        items_json = raw_items if raw_items else "[]"
    else:
        header_form = GarantiaHeaderForm()
        items_json = "[]"

    ctx = {
        "header_form": header_form,
        "produtos": produtos_qs,
        "items_json": items_json,
        "is_new": True,
    }
    return render(request, "fretes/garantias_multi_form.html", ctx)
# ---------------- Ferramentas administrativas ----------------
@permission_required('fretes.can_import_products', raise_exception=True)
def admin_import_produtos(request):
    context = {}
    if request.method == "POST" and request.FILES.get("arquivo"):
        arquivo = request.FILES["arquivo"]
        try:
            wb = load_workbook(filename=arquivo, data_only=True)
            ws = wb.active
        except Exception as e:
            messages.error(request, f"Arquivo inválido: {e}")
            return render(request, "fretes/import_produtos.html", context)

        # Normaliza cabeçalhos removendo acentos e padronizando (CM)
        import unicodedata as _ud

        def norm_header(s: str) -> str:
            if not s:
                return ""
            s = str(s).strip()
            s = _ud.normalize("NFKD", s)
            s = "".join(ch for ch in s if not _ud.combining(ch))
            s = s.upper().replace("  ", " ")
            s = s.replace("(CM)", "CM").replace("  ", " ").strip()
            return s

        headers_raw = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        headers = [norm_header(h) for h in headers_raw]
        index = {h: i for i, h in enumerate(headers)}

        # Aceita variações com/sem acentos e parênteses
        required = {
            "CODIGO": ["CODIGO"],
            "DESCRICAO": ["DESCRICAO"],
            "APLICACAO": ["APLICACAO"],
            "PESO BRUTO": ["PESO BRUTO"],
            "PESO LIQUIDO": ["PESO LIQUIDO", "PESO LIQUIDO"],
            "LARGURA CM": ["LARGURA CM"],
            "ALTURA CM": ["ALTURA CM"],
            "COMPRIMENTO CM": ["COMPRIMENTO CM"],
        }

        def find_col(key):
            for alias in required[key]:
                if alias in index:
                    return alias
            return None

        missing_keys = [k for k in required if not find_col(k)]
        if missing_keys:
            msgs = ", ".join(missing_keys)
            messages.error(request, f"Colunas obrigatórias ausentes: {msgs}")
            return render(request, "fretes/import_produtos.html", context)

            if isinstance(v, (int, float)):
                return Decimal(str(v))
            s = str(v).strip()
            s = s.replace(".", "").replace(",", ".")
            try:
                return Decimal(s)
            except InvalidOperation:
                return Decimal("0")

        criados = atualizados = linhas = 0
        try:
            with transaction.atomic():
                for row in ws.iter_rows(min_row=2):
                    linhas += 1
                    def val(colkey):
                        alias = find_col(colkey)
                        i = index.get(alias)
                        return row[i].value if i is not None else None

                    codigo = str(val("CODIGO") or "").strip()
                    if not codigo:
                        continue
                    # Concatena descrição + aplicação, normalizando quebras de linha
                    def collapse_ws(s: str) -> str:
                        return " ".join(str(s).split())
                    descricao_base = collapse_ws(val("DESCRICAO") or "")
                    aplicacao = collapse_ws(val("APLICACAO") or "")
                    descricao = (descricao_base + (f" {aplicacao}" if aplicacao else "")).strip()
                    # Respeita o tamanho do campo no modelo
                    try:
                        max_len = Produto._meta.get_field("descricao").max_length or 255
                    except Exception:
                        max_len = 255
                    if len(descricao) > max_len:
                        descricao = descricao[:max_len]

                    peso_bruto = _parse_decimal_value(val("PESO BRUTO"))
                    peso_liquido = _parse_decimal_value(val("PESO LIQUIDO"))
                    largura = _parse_decimal_value(val("LARGURA CM"))
                    altura = _parse_decimal_value(val("ALTURA CM"))
                    comprimento = _parse_decimal_value(val("COMPRIMENTO CM"))

                    _, created = Produto.objects.update_or_create(
                        codigo=codigo,
                        defaults={
                            "descricao": descricao,
                            "peso_bruto_kg": peso_bruto,
                            "peso_liquido_kg": peso_liquido,
                            "largura_cm": largura,
                            "altura_cm": altura,
                            "comprimento_cm": comprimento,
                        },
                    )
                    if created:
                        criados += 1
                    else:
                        atualizados += 1
        except Exception as e:
            messages.error(request, f"Falha ao importar na linha {linhas+1}: {e}")
            return render(request, "fretes/import_produtos.html", context)

        messages.success(request, f"Importação concluída. Linhas lidas: {linhas}. Criados: {criados}. Atualizados: {atualizados}.")

    return render(request, "fretes/import_produtos.html", context)


@permission_required('fretes.can_import_clients', raise_exception=True)
def admin_import_clientes(request):
    context = {}
    if request.method == "POST" and request.FILES.get("arquivo"):
        arquivo = request.FILES["arquivo"]
        try:
            wb = load_workbook(filename=arquivo, data_only=True)
            ws = wb.active
        except Exception as e:
            messages.error(request, f"Arquivo inválido: {e}")
            return render(request, "fretes/import_clientes.html", context)

        headers = [str(c.value).strip().upper() if c.value is not None else "" for c in next(ws.iter_rows(min_row=1, max_row=1))]
        index = {h: i for i, h in enumerate(headers)}

        required = ["NOME", "CNPJ", "ENDERECO", "CIDADE", "ESTADO", "EMAIL", "TELEFONE"]
        missing = [h for h in required if h not in index]
        if missing:
            messages.error(request, f"Colunas obrigatórias ausentes: {', '.join(missing)}")
            return render(request, "fretes/import_clientes.html", context)

        def cell(col, row):
            i = index.get(col)
            return row[i].value if i is not None else None

        def norm_str(v):
            return str(v).strip() if v is not None else ""

        def norm_cnpj(v):
            s = norm_str(v)
            return "".join(ch for ch in s if ch.isdigit())

        criados = atualizados = ignorados = 0
        with transaction.atomic():
            for row in ws.iter_rows(min_row=2):
                nome = norm_str(cell("NOME", row))
                cnpj = norm_cnpj(cell("CNPJ", row))
                if not nome or not cnpj:
                    ignorados += 1
                    continue
                endereco = norm_str(cell("ENDERECO", row))
                cidade = norm_str(cell("CIDADE", row))
                estado = norm_str(cell("ESTADO", row)).upper()[:2]
                email = norm_str(cell("EMAIL", row))
                telefone = norm_str(cell("TELEFONE", row))

                obj, created = Cliente.objects.update_or_create(
                    cnpj=cnpj,
                    defaults={
                        "nome": nome,
                        "endereco": endereco,
                        "cidade": cidade,
                        "estado": estado,
                        "email": email,
                        "telefone": telefone,
                    },
                )
                if created:
                    criados += 1
                else:
                    atualizados += 1

        messages.success(
            request,
            f"Importação de clientes concluída. Criados: {criados}. Atualizados: {atualizados}. Ignorados (faltando nome/cnpj): {ignorados}.",
        )

    return render(request, "fretes/import_clientes.html", context)


@staff_member_required
def admin_template_clientes(request):
    from io import BytesIO
    from openpyxl import Workbook
    from django.http import HttpResponse

    wb = Workbook()
    ws = wb.active
    ws.title = "Clientes"
    headers = ["NOME", "CNPJ", "ENDERECO", "CIDADE", "ESTADO", "EMAIL", "TELEFONE"]
    ws.append(headers)
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    resp = HttpResponse(
        buf.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    resp["Content-Disposition"] = 'attachment; filename="template_clientes.xlsx"'
    return resp


@permission_required('fretes.view_auditlog', raise_exception=True)
def audit_log_view(request):
    User = get_user_model()
    users = User.objects.order_by('username')
    qs = AuditLog.objects.select_related('user').all().order_by('-created_at', '-id')
    user_id = request.GET.get('user')
    action = request.GET.get('action', '').strip()
    module = request.GET.get('module', '').strip()
    start_date = request.GET.get('start_date', '').strip()
    end_date = request.GET.get('end_date', '').strip()
    q = request.GET.get('q', '').strip()

    if user_id:
        qs = qs.filter(user_id=user_id)
    if action:
        qs = qs.filter(action=action)
    if module:
        qs = qs.filter(module=module)
    if start_date:
        qs = qs.filter(created_at__date__gte=parse_date(start_date))
    if end_date:
        qs = qs.filter(created_at__date__lte=parse_date(end_date))
    if q:
        qs = qs.filter(object_repr__icontains=q) | qs.filter(path__icontains=q) | qs.filter(username__icontains=q)

    return render(request, 'fretes/auditoria.html', {
        'logs': qs[:500],
        'users': users,
        'user_id': user_id or '',
        'action': action,
        'module': module,
        'start_date': start_date,
        'end_date': end_date,
        'q': q,
    })



















@permission_required('fretes.can_use_calcular', raise_exception=True)
def calcular_view(request):
    context = {}
    if request.method == "POST":
        form = CalcularFreteForm(request.POST)
        if form.is_valid():
            pedidos = form.cleaned_data["pedidos"]
            carrier = form.cleaned_data["carrier"]
            kg_nota = form.cleaned_data["kg_nota"]
            valor_nota = form.cleaned_data.get("valor_nota")
            numero_nota = (form.cleaned_data.get("numero_nota") or "").strip()
            tipo_frete = (form.cleaned_data.get("tipo_frete") or "").strip() or None
            autorizado_por = (form.cleaned_data.get("autorizado_por") or "").strip()

            m3_total = Decimal("0")
            for p in pedidos:
                m3_p = _calcular_m3_total_pedido(p)
                p.m3 = m3_p
                p.save(update_fields=["m3"])
                m3_total += m3_p

            r = calcular_frete(m3_total, carrier, kg_nota, valor_nota)

            if len(pedidos) > 1:
                legacy_num = "MULT"
            else:
                legacy_num = pedidos[0].numero_pedido

            defaults = {
                "data_calculo": timezone.now().date(),
                "numero_nota": numero_nota,
                "valor_nota": valor_nota or 0,
                "kg_nota": kg_nota,
                "m3": r["m3"],
                "peso_cubico": r["peso_cubico"],
                "peso_usado": r["peso_usado"],
                "frete_total": r["frete_total"],
                "tipo_frete": tipo_frete or FreteCalculado.TIPO_PAGO,
                "autorizado_por": autorizado_por,
            }

            if numero_nota:
                fc, _ = FreteCalculado.objects.update_or_create(
                    numero_nota=numero_nota,
                    carrier=carrier,
                    defaults={"numero_pedido": legacy_num, **defaults},
                )
            else:
                fc = FreteCalculado.objects.create(
                    numero_pedido=legacy_num,
                    carrier=carrier,
                    **defaults,
                )
            fc.pedidos.set(pedidos)

            context.update({"resultado": r, "pedidos": pedidos, "carrier": carrier, "form": form})
            if request.headers.get("HX-Request") == "true":
                return render(request, "fretes/_resultado.html", context)
        else:
            if request.headers.get("HX-Request") == "true":
                return render(request, "fretes/_resultado.html", {"form": form})
    else:
        form = CalcularFreteForm()
    context["form"] = form
    return render(request, "fretes/calcular.html", context)

