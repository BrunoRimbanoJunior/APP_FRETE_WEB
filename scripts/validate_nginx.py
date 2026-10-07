"""Verifica o proxy real na instancia isolada, sem importar dados de negocio.

Execute dentro do web de docker-compose.validation.yml, com --access-file
apontando para o acesso.txt local. O relatorio nao inclui credenciais.
"""

import argparse
import gzip
import io
import json
import os
import re
from http.cookiejar import CookieJar
from pathlib import Path
from urllib.error import HTTPError
from urllib.parse import urlencode
from urllib.request import HTTPCookieProcessor, HTTPRedirectHandler, Request, build_opener
from zipfile import ZIP_STORED, ZipFile

import django


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--base-url", default="http://nginx")
    parser.add_argument("--access-file", required=True)
    parser.add_argument("--report", default="/tmp/nginx-functional-report.json")
    args = parser.parse_args()
    django.setup()
    from django.conf import settings
    from django.templatetags.static import static
    from django.urls import reverse
    from fretes.models import Cliente, FreteCalculado, Garantia, Pedido, Produto
    from openpyxl import Workbook, load_workbook

    if os.getenv("VALIDATION_ONLY") != "1" or settings.DATABASES["default"]["NAME"] != "validation":
        raise RuntimeError("Este teste exige o banco isolado validation e VALIDATION_ONLY=1.")

    access = Path(args.access_file).read_text(encoding="utf-8-sig")
    username = re.search(r"^Usu[aá]rio de valida[cç][aã]o: (.+)$", access, re.M).group(1).strip()
    password = re.search(r"^Senha: (.+)$", access, re.M).group(1).strip()
    base = args.base_url.rstrip("/")
    cookies = CookieJar()
    opener = build_opener(HTTPCookieProcessor(cookies))

    class NoRedirect(HTTPRedirectHandler):
        def redirect_request(self, req, fp, code, msg, headers, newurl):
            return None

    redirect_opener = build_opener(HTTPCookieProcessor(cookies), NoRedirect())
    models = (Cliente, Produto, Pedido, Garantia, FreteCalculado)

    def counts():
        return {model.__name__: model.objects.count() for model in models}

    def request(path, data=None, headers=None, follow_redirects=True):
        try:
            client = opener if follow_redirects else redirect_opener
            response = client.open(Request(base + path, data=data, headers=headers or {}), timeout=70)
        except HTTPError as error:
            response = error
        with response:
            return response.status, response.headers, response.read(), response.url

    def csrf():
        return next(cookie.value for cookie in cookies if cookie.name == "csrftoken")

    report = {"pages": {}, "uploads": {}, "assets": {}, "counts_before": counts()}
    report["redirects"] = {}
    for path, destination in (("/", "/fretes/"), ("/admin", "/admin/"), ("/fretes", "/fretes/")):
        status, headers, _, _ = request(path, headers={"Host": "localhost:8083"}, follow_redirects=False)
        assert status == 301 and headers["Location"] == destination, path
        report["redirects"][path] = destination
    status, _, html, _ = request("/admin/login/?next=/fretes/")
    assert status == 200
    token = re.search(rb'name="csrfmiddlewaretoken" value="([^"]+)"', html).group(1).decode()
    status, _, html, url = request("/admin/login/?next=/fretes/", urlencode({
        "username": username, "password": password, "csrfmiddlewaretoken": token, "next": "/fretes/"
    }).encode(), {"Content-Type": "application/x-www-form-urlencoded", "Referer": base + "/admin/login/"})
    assert status == 200 and url == base + "/fretes/", "Login falhou"
    report["login"] = "OK"
    status, _, _, _ = request("/admin/login/", headers={"X-Forwarded-Host": "invalid host", "X-Forwarded-Proto": "https"})
    assert status == 200, "Headers externos interferiram no proxy"
    report["forwarded_headers_overwritten"] = "OK"

    for name in ("index", "pedido_list", "produto_list", "cliente_list", "garantia_list",
                 "relatorios", "romaneio", "garantias_gerencial", "audit_logs"):
        status, _, html, _ = request(reverse("fretes:" + name))
        assert status == 200 and b"<main" in html, name
        report["pages"][name] = status
    for name, model in (("pedido_update", Pedido), ("produto_update", Produto),
                        ("cliente_update", Cliente), ("garantia_update", Garantia)):
        obj = model.objects.order_by("pk").first()
        assert obj is not None, name
        status, _, html, _ = request(reverse("fretes:" + name, args=[obj.pk]))
        assert status == 200 and b"<main" in html, name
        report["pages"][name] = status
    code = Produto.objects.order_by("pk").first().codigo
    status, _, html, _ = request(reverse("fretes:produto_list") + "?" + urlencode({"q": code}))
    assert status == 200 and code.encode() in html
    report["filter"] = "OK"

    # XLSX valido com entrada extra nao referenciada: >1 MB, sem cabecalhos de importacao.
    # O Django deve ler a planilha e recusar os campos, sem criar/alterar cadastros.
    workbook = Workbook()
    workbook.active.append(["VALIDACAO SEM DADOS"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    with ZipFile(buffer, "a", compression=ZIP_STORED) as archive:
        archive.writestr("validation-padding.bin", b"x" * (2 * 1024 * 1024))
    payload = buffer.getvalue()
    boundary = "nginx-validation-boundary"

    def multipart(content):
        return (
            f"--{boundary}\r\nContent-Disposition: form-data; name=\"arquivo\"; filename=\"validation.xlsx\"\r\n"
            "Content-Type: application/vnd.openxmlformats-officedocument.spreadsheetml.sheet\r\n\r\n"
        ).encode() + content + f"\r\n--{boundary}--\r\n".encode()

    for name in ("admin_import_produtos", "admin_import_clientes"):
        path = reverse("fretes:" + name)
        status, _, html, _ = request(path, multipart(payload), {
            "Content-Type": f"multipart/form-data; boundary={boundary}",
            "X-CSRFToken": csrf(), "Referer": base + path,
        })
        assert status == 200 and "Colunas obrigatórias ausentes" in html.decode(), name
        report["uploads"][name] = {"bytes": len(payload), "status": status, "django_validation": "OK"}
        status, _, _, _ = request(path, b"x" * (21 * 1024 * 1024), {"Content-Type": "application/octet-stream"})
        assert status == 413, name
        report["uploads"][name]["above_20mb"] = status
    status, _, _, _ = request("/fretes/calcular/", b"x" * (2 * 1024 * 1024), {"Content-Type": "application/octet-stream"})
    assert status == 413
    report["normal_route_above_1mb"] = status

    report["exports"] = {}
    for extension in ("pdf", "xlsx"):
        status, _, content, _ = request(reverse("fretes:relatorios") + "?export=" + extension)
        assert status == 200
        if extension == "pdf":
            assert content.startswith(b"%PDF")
        else:
            exported = load_workbook(io.BytesIO(content), read_only=True)
            assert exported.active.max_row >= 1
            if report["counts_before"]["FreteCalculado"]:
                assert exported.active.max_row > 1
            exported.close()
        report["exports"][extension] = {"status": status, "bytes": len(content)}

    status, headers, compressed, _ = request("/fretes/", headers={"Accept-Encoding": "gzip"})
    assert status == 200 and headers["Content-Encoding"] == "gzip" and b"<main" in gzip.decompress(compressed)
    report["html_gzip"] = "OK"
    import brotli
    for asset in ("vendor/bootstrap/bootstrap-5.3.8.min.css", "vendor/htmx/htmx-2.0.11.min.js", "fretes/app.css"):
        path = static(asset)
        status, _, original, _ = request(path, headers={"Accept-Encoding": "identity"})
        assert status == 200
        status, headers, content, _ = request(path, headers={"Accept-Encoding": "br"})
        assert status == 200 and headers["Content-Encoding"] == "br"
        assert brotli.decompress(content) == original and "immutable" in headers["Cache-Control"]
        report["assets"][asset] = {"original_bytes": len(original), "brotli_bytes": len(content), "cache": "immutable"}
    if os.getenv("CHECK_NGINX_HEALTH") == "1":
        status, headers, content, _ = request("/nginx-health")
        assert status == 200 and content == b"ok\n" and "nginx/" not in headers["Server"]
        report["nginx_health"] = "OK"
    report["counts_after"] = counts()
    assert report["counts_before"] == report["counts_after"], "Dados de negocio alterados"
    Path(args.report).write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
